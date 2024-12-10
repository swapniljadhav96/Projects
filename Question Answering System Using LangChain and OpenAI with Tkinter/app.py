import os
import tkinter as tk
from tkinter import scrolledtext
from tkinter import filedialog, messagebox
from docx import Document
from PyPDF2 import PdfReader
from langchain import LLMChain, PromptTemplate
from langchain.llms import OpenAI

# Set up the OpenAI API key in the environment
os.environ['OPENAI_API_KEY'] = "Enter Your OpenAI Key here"

# Function to read PDF files
def read_pdf_file(file_path):
    pdf_text = ""
    reader = PdfReader(file_path)
    for page in reader.pages:
        pdf_text += page.extract_text()
    return pdf_text

# Function to read Word files
def read_word_file(file_path):
    doc = Document(file_path)
    word_text = ""
    for paragraph in doc.paragraphs:
        word_text += paragraph.text + "\n"
    return word_text

# Function to answer questions based on the provided context
def answer_question(question, context):
    return qa_chain.run({"question": question, "context": context})

# Load all text from files
def load_files():
    pdf1 = read_pdf_file('India A Comprehensive Overview.pdf')
    pdf2 = read_pdf_file('Indias Diverse States and Territories.pdf')
    word1 = read_word_file('Indias_Education_Healthcare_and_Social_Development.docx')
    word2 = read_word_file('documentsIndias_Natural_Beauty_and_Wildlife.docx')
    return word1 + "\n" + word2 + "\n" + pdf1 + "\n" + pdf2

# Configure LangChain and OpenAI
prompt_template = PromptTemplate(              # define how the input prompts to the language model should be structured
    input_variables=["question", "context"],   #placeholders
    template="Based on the following context, please provide a concise and accurate answer to the question.\n\nContext: {context}\n\nQuestion: {question}\nAnswer:"
)

llm = OpenAI(temperature=0.7)  # initialize the OpenAI language model.
qa_chain = LLMChain(llm=llm, prompt=prompt_template)

# Initialize Tkinter
root = tk.Tk()
root.title("Question Answering System")

# Create TextBox for Question input
question_label = tk.Label(root, text="Enter your question:")
question_label.pack(pady=5)
question_entry = tk.Entry(root, width=80)
question_entry.pack(pady=5)

# Create ScrolledText for displaying the answer
answer_label = tk.Label(root, text="Answer:")
answer_label.pack(pady=5)
answer_text = scrolledtext.ScrolledText(root, wrap=tk.WORD, width=80, height=10)
answer_text.pack(pady=5)

# Load the document text once
all_text = load_files()

# Function to process the question and display the answer
def get_answer():
    question = question_entry.get()
    if not question:
        messagebox.showwarning("Input Error", "Please enter a question.")
        return

    try:
        answer = answer_question(question, all_text)
        answer_text.delete(1.0, tk.END)
        answer_text.insert(tk.END, answer)
    except Exception as e:
        messagebox.showerror("Error", str(e))

# Create Button to submit the question
submit_button = tk.Button(root, text="Get Answer", command=get_answer)
submit_button.pack(pady=20)

# Run the Tkinter main loop
root.mainloop()